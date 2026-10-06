"""Typeset the paper as an IEEE-style two-column Word document (and PDF via Word).

The LaTeX sources in this folder stay the single source of truth; this script
converts the small LaTeX subset they use (sections, emphasis, citations,
cross-references, lists, inline math, equations, figures, booktabs tables) so a
PDF can be produced on a machine without a TeX installation.

    python build_docx.py            -> SmartSwap_paper.docx (+ .pdf if Word is installed)
"""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
BUILD = HERE / "build"; BUILD.mkdir(exist_ok=True)
FONT = "Times New Roman"

# ---------------------------------------------------------------- source
main = (HERE / "main.tex").read_text(encoding="utf-8")
order = re.findall(r"\\input\{(sections_[^}]+)\}", main)
source = "\n".join((HERE / f"{name}.tex").read_text(encoding="utf-8") for name in order)
source = re.sub(r"(?<!\\)%.*", "", source)  # comments (incl. %%HYBRID%% placeholders)
title = re.search(r"\\title\{(.+?)\}\n", main, re.S).group(1)
authors = re.search(r"\\IEEEauthorblockN\{(.+?)\}", main).group(1)
affil = re.search(r"\\IEEEauthorblockA\{(.+?)\}\}", main, re.S).group(1).replace("\\\\", "\n")
keywords = re.search(r"\\begin\{IEEEkeywords\}(.+?)\\end\{IEEEkeywords\}", main, re.S).group(1).strip()
artifact = re.search(r"\\section\*\{Artifact\}\n(.+?)\n\n", main, re.S).group(1)

# ---------------------------------------------------------------- numbering
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]
labels: dict[str, str] = {}
sec = sub = fig = tab = eq = 0
for m in re.finditer(r"\\(section|subsection)\{[^}]*\}(\s*\\label\{([^}]+)\})?|\\begin\{(figure\*?|table\*?|equation)\}(.*?)\\end\{\4\}", source, re.S):
    if m.group(1) == "section": sec += 1; sub = 0; key, val = m.group(3), ROMAN[sec - 1]
    elif m.group(1) == "subsection": sub += 1; key, val = m.group(3), f"{ROMAN[sec - 1]}-{chr(64 + sub)}"
    else:
        kind, body = m.group(4), m.group(5)
        lab = re.search(r"\\label\{([^}]+)\}", body); key = lab.group(1) if lab else None
        if kind.startswith("figure"): fig += 1; val = str(fig)
        elif kind.startswith("table"): tab += 1; val = ROMAN[tab - 1]
        else: eq += 1; val = str(eq)
    if key: labels[key] = val

bib_text = (HERE / "refs.bib").read_text(encoding="utf-8")
cite_order: list[str] = []
for group in re.findall(r"\\cite\{([^}]+)\}", source):
    for key in (k.strip() for k in group.split(",")):
        if key not in cite_order: cite_order.append(key)

# ---------------------------------------------------------------- inline LaTeX -> runs
SYMBOLS = {r"\rho": "ρ", r"\times": "×", r"\geq": "≥", r"\ge": "≥", r"\leq": "≤", r"\rightarrow": "→", r"\theta": "θ", r"\ell": "ℓ",
           r"\cdot": "·", r"\langle": "⟨", r"\rangle": "⟩", r"\in": "∈", r"\hat": "", r"\mathrm": "", r"\text": "", r"\,": "\u2009", r"\;": " ", r"\!": "", r"\ ": " ", r"\%": "%"}


def math_to_text(expr: str) -> str:
    expr = re.sub(r"\\(?:mathrm|text)\{([^}]*)\}", r"\1", expr)
    expr = re.sub(r"\\hat\{([^}]*)\}", lambda m: m.group(1) + "\u0302", expr)
    for k in sorted(SYMBOLS, key=len, reverse=True): expr = expr.replace(k, SYMBOLS[k])
    expr = re.sub(r"_\{([^}]*)\}|_(\w)", lambda m: "\x03" + (m.group(1) or m.group(2)) + "\x04", expr)  # real subscript run
    expr = re.sub(r"\^\{([^}]*)\}|\^(\w+)", lambda m: "^" + (m.group(1) or m.group(2)), expr)
    expr = re.sub(r"\^(\w+)", lambda m: m.group(1).translate(str.maketrans("0123456789", "⁰¹²³⁴⁵⁶⁷⁸⁹")) if m.group(1).isdigit() else "^" + m.group(1), expr)
    return expr.replace("{", "").replace("}", "")


def plain(text: str) -> str:
    """Text-level substitutions applied to every run."""
    text = text.replace("---", "\u2014").replace("--", "\u2013").replace("``", "\u201c").replace("''", "\u201d").replace("`", "\u2018")
    text = text.replace("\\ ", " ").replace("~", "\u00a0").replace(r"\,", "\u2009").replace(r"\%", "%").replace(r"\&", "&").replace(r"\_", "_").replace(r"\S", "§").replace(r"\eg", "e.g.,")
    text = re.sub(r"\\textbackslash", "\\\\", text)
    return re.sub(r"\s+", " ", text)


def resolve(text: str) -> str:
    text = re.sub(r"\\cite\{([^}]+)\}", lambda m: "[" + ", ".join(str(cite_order.index(k.strip()) + 1) for k in m.group(1).split(",")) + "]", text)
    text = re.sub(r"\\ref\{([^}]+)\}", lambda m: labels.get(m.group(1), "??"), text)
    text = re.sub(r"\\label\{[^}]*\}", "", text)
    text = re.sub(r"\$([^$]+)\$", lambda m: "\x01" + math_to_text(m.group(1)) + "\x02", text)  # math -> italic-marked
    return text


def runs_of(text: str, style: tuple[bool, bool, bool] = (False, False, False)) -> list[tuple[str, bool, bool, bool]]:
    """Split LaTeX text into (text, bold, italic, mono) runs; handles nesting."""
    out: list[tuple[str, bool, bool, bool]] = []
    pattern = re.compile(r"\\(textbf|emph|textit|texttt|url)\{")
    pos = 0
    while True:
        m = pattern.search(text, pos)
        if not m: out.append((text[pos:], *style)); break
        out.append((text[pos:m.start()], *style))
        depth, i = 1, m.end()
        while depth and i < len(text):
            depth += {"{": 1, "}": -1}.get(text[i], 0); i += 1
        inner = text[m.end():i - 1]
        b, it, mono = style
        cmd = m.group(1)
        out += runs_of(inner, (b or cmd == "textbf", it or cmd in ("emph", "textit"), mono or cmd in ("texttt", "url")))
        pos = i
    final = []
    for t, b, it, mono in out:
        # \x01..\x02 marks math (italic); \x03..\x04 a subscript inside it (prefixed \x05 for add_run)
        state = {"math": False, "sub": False}; buf = []

        def flush() -> None:
            if buf:
                text = "".join(buf)
                final.append((("\x05" if state["sub"] else "") + (text if state["math"] else plain(text)), b, it or state["math"], mono)); buf.clear()
        for ch in t:
            if ch in "\x01\x02\x03\x04":
                flush()
                if ch in "\x01\x02": state["math"] = ch == "\x01"
                else: state["sub"] = ch == "\x03"
            else: buf.append(ch)
        flush()
    return final


# ---------------------------------------------------------------- document helpers
doc = Document()
st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
pf = st.paragraph_format; pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0


def set_cols(section, n: int) -> None:
    cols = section._sectPr.xpath("./w:cols")
    el = cols[0] if cols else OxmlElement("w:cols")
    el.set(qn("w:num"), str(n)); el.set(qn("w:space"), "360")
    if not cols: section._sectPr.append(el)


def page_setup(section) -> None:
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(0.75); section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.625); section.right_margin = Inches(0.625)


def add_run(p, text: str, size: float, bold: bool = False, italic: bool = False, mono: bool = False):
    sub = text.startswith("\x05"); text = text.lstrip("\x05")
    r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.subscript = sub
    r.font.name = "Consolas" if mono else FONT
    return r


def para(text: str = "", align=WD_ALIGN_PARAGRAPH.JUSTIFY, size: float = 10, indent: bool = True, space_after: float = 0, runs=None, italic=False):
    p = doc.add_paragraph(); p.alignment = align
    p.paragraph_format.first_line_indent = Inches(0.14) if indent else None
    p.paragraph_format.space_after = Pt(space_after)
    for t, b, it, mono in (runs if runs is not None else runs_of(resolve(text))):
        add_run(p, t, size - 0.5 if mono else size, b, it or italic, mono)
    return p


def heading(text: str, level: int) -> None:
    if level == 1:
        global sec_no
        sec_no += 1
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(4); p.paragraph_format.keep_with_next = True
        r = p.add_run(f"{ROMAN[sec_no - 1]}. " if sec_no else ""); r.font.size = Pt(10)
        r = p.add_run(plain(text).upper()); r.font.size = Pt(10); r.font.small_caps = True
    elif level == 0:  # unnumbered (Artifact, References)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(4); p.paragraph_format.keep_with_next = True
        r = p.add_run(plain(text).upper()); r.font.size = Pt(10); r.font.small_caps = True
    else:
        global sub_no
        sub_no += 1
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2); p.paragraph_format.keep_with_next = True
        r = p.add_run(f"{chr(64 + sub_no)}. {plain(text)}"); r.italic = True; r.font.size = Pt(10)


def equation_image(expr: str, number: str) -> None:
    expr = re.sub(r"\\label\{[^}]*\}", "", expr).strip().rstrip(".").rstrip(",")
    expr = re.sub(r"\\text\{([^}]*)\}", lambda m: r"\mathrm{" + m.group(1).strip().replace(" ", r"\ ") + "}", expr)
    expr = expr.replace(r"\!", "").replace(r"\;", r"\ ").replace(r"\ge", r"\geq").replace(r"\geqq", r"\geq")
    path = BUILD / f"eq{number}.png"
    f = plt.figure(figsize=(0.01, 0.01)); f.text(0, 0, f"${expr}$", fontsize=11, family="serif")
    f.savefig(path, dpi=400, bbox_inches="tight", pad_inches=0.03, transparent=False); plt.close(f)
    t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    c0, c1 = t.rows[0].cells; c0.width = Inches(3.0); c1.width = Inches(0.4)
    p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER; p0.add_run().add_picture(str(path), height=Inches(0.36 if r"\frac" in expr else 0.17))
    p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p1.add_run(f"({number})").font.size = Pt(10)
    p1.paragraph_format.space_before = Pt(6)


def caption(text: str, kind: str, number: str) -> None:
    if kind == "Fig.":
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.space_after = Pt(8)
        r = p.add_run(f"Fig. {number}. "); r.font.size = Pt(8)
        for t, b, it, mono in runs_of(resolve(text)):
            add_run(p, t, 8, b, it, mono)
    else:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(6); p.paragraph_format.keep_with_next = True
        r = p.add_run(f"TABLE {number}"); r.font.size = Pt(8)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(3); p.paragraph_format.keep_with_next = True
        for t, b, it, mono in runs_of(resolve(text)):
            add_run(p, t, 8, b, it, mono).font.small_caps = not it


def border(cell, side: str, size: int) -> None:
    tcPr = cell._tc.get_or_add_tcPr(); borders = tcPr.find(qn("w:tcBorders"))
    if borders is None: borders = OxmlElement("w:tcBorders"); tcPr.append(borders)
    el = OxmlElement(f"w:{side}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(size)); el.set(qn("w:color"), "000000"); borders.append(el)


def booktabs_table(tex: str, width_in: float) -> None:
    body = tex.split(r"\toprule")[1].split(r"\bottomrule")[0]
    rows, rule_after = [], set()
    for chunk in re.split(r"\\\\", body):
        chunk = chunk.strip()
        if not chunk: continue
        if chunk.startswith(r"\midrule"):
            rule_after.add(len(rows) - 1); chunk = chunk[len(r"\midrule"):].strip()
            if not chunk: continue
        cells = [c.strip() for c in re.split(r"(?<!\\)&", chunk)]
        rows.append(cells)
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    weights = {6: [1.2, 1.9, 0.7, 1.5, 0.7, 0.9], 5: [0.8, 1.1, 1.4, 0.6, 0.6], 4: [1.0, 0.9, 0.8, 0.6]}.get(ncols, [1.0] * ncols)
    for j, w in enumerate(weights):
        for cell in t.columns[j].cells: cell.width = Inches(w * width_in / sum(weights))
    for i, cells in enumerate(rows):
        if len(cells) == 2 and ncols > 2 and cells[1].startswith(r"\multicolumn"):  # "label & \multicolumn{3}{r}{text}"
            mc = re.match(r"\\multicolumn\{\d+\}\{\w\}\{(.*)\}$", cells[1]); merged = t.cell(i, 1).merge(t.cell(i, ncols - 1)); cells = [cells[0], mc.group(1) if mc else cells[1]]
        for j, text in enumerate(cells):
            cell = t.cell(i, j); p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j < (2 if ncols >= 5 else 1) else WD_ALIGN_PARAGRAPH.RIGHT
            for tt, b, it, mono in runs_of(resolve(text)):
                add_run(p, tt, 7.5, b or i == 0, it, mono)
            if i == 0: border(cell, "top", 8); border(cell, "bottom", 4)
            if i in rule_after: border(cell, "bottom", 4)
            if i == len(rows) - 1: border(cell, "bottom", 8)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def figure(path: Path, width_in: float) -> None:
    png = path.with_suffix(".png")
    if not png.exists():  # the PDFs are vector; Word needs a raster
        alt = HERE.parent / "research" / "results"
        hits = list(alt.rglob(png.name)); png = hits[0] if hits else png
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(4); p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(png), width=Inches(width_in))


def new_section(cols: int) -> None:
    s = doc.add_section(WD_SECTION.CONTINUOUS); page_setup(s); set_cols(s, cols)


# ---------------------------------------------------------------- build
page_setup(doc.sections[0]); set_cols(doc.sections[0], 1)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(10)
r = p.add_run(plain(title)); r.font.size = Pt(22)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run(plain(authors)); r.font.size = Pt(11)
for line in affil.split("\n"):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run(plain(line.strip())); r.font.size = Pt(9); r.italic = "@" not in line
doc.add_paragraph().paragraph_format.space_after = Pt(6)
new_section(2)

sec_no = 0; sub_no = 0
abstract = re.search(r"\\begin\{abstract\}(.+?)\\end\{abstract\}", source, re.S).group(1).strip()
p = para(runs=[("Abstract\u2014", True, True, False)] + [(t, True, it, m) for t, b, it, m in runs_of(resolve(abstract))], size=9, indent=False, space_after=6)
p = para(runs=[("Index Terms\u2014", True, True, False), (plain(keywords) + ".", True, False, False)], size=9, indent=False, space_after=6)

body = source[source.index(r"\end{abstract}") + len(r"\end{abstract}"):]
blocks = re.split(r"\n\s*\n", body)
for block in blocks:
    block = block.strip()
    if not block: continue
    # split a block that mixes headings / environments with text
    tokens = re.split(r"(\\section\{[^}]*\}(?:\s*\\label\{[^}]*\})?|\\subsection\{[^}]*\}(?:\s*\\label\{[^}]*\})?|\\begin\{(?:figure\*?|table\*?|equation|enumerate|itemize)\}.*?\\end\{(?:figure\*?|table\*?|equation|enumerate|itemize)\})", block, flags=re.S)
    for tok in tokens:
        tok = tok.strip()
        if not tok: continue
        if tok.startswith(r"\section{"):
            heading(re.match(r"\\section\{([^}]*)\}", tok).group(1), 1); sub_no = 0
        elif tok.startswith(r"\subsection{"):
            heading(re.match(r"\\subsection\{([^}]*)\}", tok).group(1), 2)
        elif tok.startswith(r"\begin{equation}"):
            inner = tok[len(r"\begin{equation}"):-len(r"\end{equation}")]
            lab = re.search(r"\\label\{([^}]+)\}", inner); equation_image(inner, labels.get(lab.group(1), "?") if lab else str(len(labels)))
        elif tok.startswith((r"\begin{enumerate}", r"\begin{itemize}")):
            numbered = tok.startswith(r"\begin{enumerate}")
            items = [i.strip() for i in re.split(r"\\item", re.sub(r"\\(begin|end)\{(enumerate|itemize)\}", "", tok)) if i.strip()]
            for n, item in enumerate(items, 1):
                p = para(runs=[(f"{n}) " if numbered else "\u2022 ", False, False, False)] + runs_of(resolve(item)), indent=False, space_after=2)
                p.paragraph_format.left_indent = Inches(0.18); p.paragraph_format.first_line_indent = Inches(-0.16)
        elif tok.startswith(r"\begin{figure"):
            wide = tok.startswith(r"\begin{figure*}")
            img = re.search(r"\\includegraphics(?:\[[^]]*\])?\{([^}]+)\}", tok).group(1)
            cap = re.search(r"\\caption\{(.+)\}\s*\\label", tok, re.S).group(1)
            lab = re.search(r"\\label\{([^}]+)\}", tok).group(1)
            if wide: new_section(1)
            figure(HERE / img, 6.9 if wide else 3.3); caption(cap, "Fig.", labels[lab])
            if wide: new_section(2)
        elif tok.startswith(r"\begin{table"):
            wide = tok.startswith(r"\begin{table*}")
            cap = re.search(r"\\caption\{(.+?)\}\s*\\label", tok, re.S).group(1)
            lab = re.search(r"\\label\{([^}]+)\}", tok).group(1)
            src = re.search(r"\\input\{([^}]+)\}", tok).group(1)
            if wide: new_section(1)
            caption(cap, "TABLE", labels[lab]); booktabs_table((HERE / src).read_text(encoding="utf-8"), 6.9 if wide else 3.3)
            if wide: new_section(2)
        else:
            text = " ".join(tok.split())
            m = re.match(r"\\textbf\{([^}]*)\}\s*(.*)", text)
            para(text)

heading("Artifact", 0); para(artifact.replace(r"\url{https://github.com/ANONYMIZED}", "https://github.com/ANONYMIZED"))

# ---------------------------------------------------------------- references (IEEE style)
def bib_entries() -> dict[str, dict[str, str]]:
    out = {}
    for m in re.finditer(r"@(\w+)\{([^,]+),(.*?)\n\}", bib_text, re.S):
        fields = dict((k.lower(), v) for k, v in re.findall(r"(\w+)\s*=\s*\{((?:[^{}]|\{[^{}]*\})*)\}", m.group(3)))
        fields["type"] = m.group(1).lower(); out[m.group(2).strip()] = fields
    return out


def clean(s: str) -> str:
    s = re.sub(r"\\url\{([^}]*)\}", r"\1", s)
    s = s.replace(r"{\'e}", "é").replace(r"{\'\i}", "í").replace(r"\'\i", "í").replace(r"\&", "&").replace(r"\_", "_").replace("--", "\u2013")
    return re.sub(r"[{}]", "", s)


def initials(name: str) -> str:
    name = clean(name).strip()
    if "," in name: last, first = [x.strip() for x in name.split(",", 1)]
    else: parts = name.split(); first, last = " ".join(parts[:-1]), parts[-1]
    return " ".join(p[0] + "." for p in re.split(r"[\s]+", first) if p) + " " + last if first else last


entries = bib_entries()
heading("References", 0)
for n, key in enumerate(cite_order, 1):
    e = entries[key]; authors_ = [initials(a) for a in e.get("author", "").split(" and ")]
    who = authors_[0] if len(authors_) == 1 else ", ".join(authors_[:-1]) + (", and " if len(authors_) > 2 else " and ") + authors_[-1]
    parts: list[tuple[str, bool, bool, bool]] = [(f"[{n}]\u2003{who}, ", False, False, False)]
    if e["type"] == "book":
        parts += [(clean(e["title"]), False, True, False), (f", {e.get('edition', '') + ' ed., ' if e.get('edition') else ''}{clean(e.get('publisher', ''))}, {e['year']}.", False, False, False)]
    else:
        parts.append((f"\u201c{clean(e['title'])},\u201d ", False, False, False))
        if e["type"] == "inproceedings": parts += [("in ", False, False, False), (clean(e["booktitle"]), False, True, False), (f", {e['year']}.", False, False, False)]
        elif e["type"] == "article":
            parts += [(clean(e["journal"]), False, True, False), (f", vol. {e.get('volume', '')}" + (f", no. {e['number']}" if e.get("number") else "") + (f", pp. {clean(e['pages'])}" if e.get("pages") else "") + f", {e['year']}.", False, False, False)]
        else: parts.append((f"{clean(e.get('howpublished', ''))}, {e.get('year', '')}{'. ' + clean(e['note']) if e.get('note') else ''}.".replace(", .", "."), False, False, False))
    p = para(runs=parts, size=8, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    p.paragraph_format.left_indent = Inches(0.3); p.paragraph_format.first_line_indent = Inches(-0.3)

out = HERE / "SmartSwap_paper.docx"; doc.save(out); print("wrote", out)

if "--no-pdf" not in sys.argv and sys.platform == "win32":
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{out}'); $d.SaveAs2('{out.with_suffix('.pdf')}', 17); $d.Close(0); $w.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True, timeout=180)
    print("wrote", out.with_suffix(".pdf"))
